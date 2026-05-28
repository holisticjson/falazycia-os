"""Konwertuje docx/xlsx z folderu Jan Szopa do MD"""
import docx, openpyxl
from pathlib import Path

base = Path(r"G:\Mój dysk\Kursy Szkolenia Marketing WWW A.I. APLIKACJE\Jan Szopa - Akademia Zdalnej Agencji Marketingowej")
out = Path(r"c:\Aplikacje MVP\Holistic Jason\deploy\knowledge")

for f in base.rglob("*.docx"):
    doc = docx.Document(str(f))
    md = f"# {f.stem}\n\n"
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.startswith("Heading"):
            md += f"## {txt}\n\n"
        else:
            md += f"{txt}\n\n"
    outf = out / f"{f.stem}.md"
    outf.write_text(md, encoding="utf-8")
    print(f"DOCX OK: {f.name}")

for f in base.rglob("*.xlsx"):
    wb = openpyxl.load_workbook(str(f), read_only=True, data_only=True)
    md = f"# {f.stem}\n\n"
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        h = [str(c) if c else "" for c in rows[0]]
        md += "| " + " | ".join(h) + " |\n"
        md += "| " + " | ".join(["---"] * len(h)) + " |\n"
        for row in rows[1:]:
            cells = [str(c).replace("|", "").replace("\n", " ") if c else "" for c in row]
            while len(cells) < len(h):
                cells.append("")
            md += "| " + " | ".join(cells[:len(h)]) + " |\n"
    wb.close()
    outf = out / f"{f.stem}.md"
    outf.write_text(md, encoding="utf-8")
    print(f"XLSX OK: {f.name}")

gdocs = list(base.rglob("*.gdoc"))
print(f"\n{len(gdocs)} plikow .gdoc (wymagaja recznego pobrania jako .docx z Google Drive):")
for g in gdocs:
    print(f"  - {g.name}")
